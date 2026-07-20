# -*- coding: utf-8 -*-
"""生成交付 Excel：Sheet1 前后对比 / Sheet2 全字段 / Sheet3 标签调整建议表 / Sheet4 非银发清单 / Sheet5 字段字典。
用法: python _excel.py
"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = "G:/workbuddy/2026-06-28-23-34-20/silver-pulse"
RUN = os.path.join(BASE, "scores_draft/run")
DB_PATH = os.path.join(BASE, "data/enterprise/all_enterprises.json")

db = json.load(open(DB_PATH, encoding="utf-8"))
by = {e["serial"]: e for e in db}
before = json.load(open(os.path.join(RUN, "before_full.json"), encoding="utf-8")) if os.path.exists(os.path.join(RUN,"before_full.json")) else {}
before_by = {e["serial"]: e for e in before}
proc = json.load(open(os.path.join(RUN, "processed.json"), encoding="utf-8")) if os.path.exists(os.path.join(RUN,"processed.json")) else []
proc = [s for s in proc if s in by]
tag_review = json.load(open(os.path.join(RUN, "tag_review_all.json"), encoding="utf-8")) if os.path.exists(os.path.join(RUN,"tag_review_all.json")) else []
nonsilver = json.load(open(os.path.join(RUN, "nonsilver_all.json"), encoding="utf-8")) if os.path.exists(os.path.join(RUN,"nonsilver_all.json")) else []

def jdump(v):
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    return "" if v is None else str(v)

def name_of(e):
    n = e.get("name","") or ""; nc = e.get("name_cn","") or ""
    if nc and nc != n: return f"{n}（{nc}）" if n else nc
    return n

hdr_fill = PatternFill("solid", fgColor="1F4E78")
hdr_font = Font(bold=True, color="FFFFFF", size=10)
green_fill = PatternFill("solid", fgColor="E2EFDA")
wrap = Alignment(wrap_text=True, vertical="top")
ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()
# ---------- Sheet1 前后对比 ----------
ws1 = wb.active; ws1.title = "Sheet1_前后对比"
h1 = ["编号","企业名称","一级标签","银发判定","信号强度(前)","信号强度(后)",
      "研究价值(前)","研究价值(后)","信息量","差异化","可复制",
      "一句话定位(后)","推荐理由v1(后)","推荐理由v2(后)","推荐理由v3(后)"]
ws1.append(h1)
for s in proc:
    a = by[s]; b = before_by.get(s, {})
    rec = a.get("recommend") or {}
    if isinstance(rec, str): rec = {}
    row = [s, name_of(a),
           "、".join(a.get("tag_l1") or []) or "、".join(b.get("tag_l1") or []),
           a.get("silver_verdict") or b.get("silver_verdict") or "",
           jdump(b.get("signal_strength")), jdump(a.get("signal_strength")),
           jdump(b.get("research_value")), jdump(a.get("research_value")),
           jdump(rec.get("info_score")), jdump(rec.get("diff_score")), jdump(rec.get("copy_score")),
           jdump(a.get("desc_cn")), jdump(rec.get("rec_v1")), jdump(rec.get("rec_v2")), jdump(rec.get("rec_v3"))]
    ws1.append(row)
for c in range(1, len(h1)+1):
    ws1.cell(1,c).fill=hdr_fill; ws1.cell(1,c).font=hdr_font; ws1.cell(1,c).alignment=ctr
widths1=[9,26,14,11,11,11,11,11,7,7,7,40,46,46,46]
for i,w in enumerate(widths1,1): ws1.column_dimensions[get_column_letter(i)].width=w
for r in range(2, ws1.max_row+1):
    for c in [12,13,14,15]: ws1.cell(r,c).alignment=wrap
    for c in [6,8]: ws1.cell(r,c).fill=green_fill
ws1.freeze_panes="A2"

# ---------- Sheet2 全字段 ----------
ws2 = wb.create_sheet("Sheet2_全字段")
preferred = ["serial","name","name_cn","region","founded","hq","stage","tag_l1","tag_l2",
             "business_tags","category_l1","category_l2","description","desc_cn","highlights",
             "events","funding_latest","funding_total","investors","payor_model","news_coverage",
             "source","website_url","crunchbase_url","business_model","business_model_cn","tags",
             "signal_strength","info_score","diff_score","copy_score","research_value","recommend",
             "silver_verdict","silver_reason","update_time","value_score","needs_review","ingest_time"]
allkeys=set()
for s in proc: allkeys.update(by[s].keys())
extra=[k for k in allkeys if k not in preferred]
cols2=preferred+sorted(extra)
ws2.append(cols2)
for s in proc:
    e=by[s]; ws2.append([jdump(e.get(k)) for k in cols2])
for c in range(1,len(cols2)+1):
    ws2.cell(1,c).fill=hdr_fill; ws2.cell(1,c).font=hdr_font; ws2.cell(1,c).alignment=ctr
    ws2.column_dimensions[get_column_letter(c)].width=22
for ci,k in enumerate(cols2,1):
    if k in ("desc_cn","description","highlights","events","recommend","news_coverage","business_tags"):
        ws2.column_dimensions[get_column_letter(ci)].width=50
ws2.freeze_panes="B2"

# ---------- Sheet3 标签调整建议表 ----------
ws3 = wb.create_sheet("Sheet3_标签调整建议")
h3=["编号","企业名称","企业介绍","原一级标签","原二级标签","原business_tags","建议操作(增/删/改 + 标签 + 理由)"]
ws3.append(h3)
tr_by={t["serial"]:t for t in tag_review}
for s in proc:
    t=tr_by.get(s)
    if not t: continue
    a=by[s]
    old=t.get("old_tags") or {}
    sug=t.get("suggested") or []
    sug_txt="\n".join(f"[{x.get('action')}] {x.get('tag')}：{x.get('reason')}" for x in sug) if sug else "（无建议）"
    ws3.append([s, name_of(a), t.get("intro",""),
                "、".join(old.get("tag_l1") or a.get("tag_l1") or []),
                "、".join(old.get("tag_l2") or a.get("tag_l2") or []),
                jdump(old.get("business_tags") or a.get("business_tags") or {}),
                sug_txt])
for c in range(1,len(h3)+1):
    ws3.cell(1,c).fill=hdr_fill; ws3.cell(1,c).font=hdr_font; ws3.cell(1,c).alignment=ctr
for i,w in enumerate([9,24,46,16,18,30,60],1): ws3.column_dimensions[get_column_letter(i)].width=w
for r in range(2, ws3.max_row+1):
    for c in [3,7]: ws3.cell(r,c).alignment=wrap
ws3.freeze_panes="A2"

# ---------- Sheet4 非银发清单 ----------
ws4 = wb.create_sheet("Sheet4_非银发清单")
ws4.append(["编号","企业名称","判定","理由"])
for x in nonsilver:
    a=by.get(x.get("serial"),{})
    ws4.append([x.get("serial"), name_of(a) or x.get("name"), x.get("verdict"), x.get("reason")])
for c in range(1,5):
    ws4.cell(1,c).fill=hdr_fill; ws4.cell(1,c).font=hdr_font; ws4.cell(1,c).alignment=ctr
for i,w in enumerate([9,24,14,70],1): ws4.column_dimensions[get_column_letter(i)].width=w
for r in range(2, ws4.max_row+1): ws4.cell(r,4).alignment=wrap
ws4.freeze_panes="A2"

# ---------- Sheet5 字段字典 ----------
ws5 = wb.create_sheet("Sheet5_字段字典")
FIELD_DICT=[
 ("serial","企业内部编号","主键，不动"),
 ("name/name_cn","英文名/中文名","name 不动；name_cn 可补"),
 ("region","地区 国内/海外","影响可复制判断"),
 ("founded","成立年份","本轮核实纠正"),
 ("stage","成长阶段","已上市/被收购/融资中/未披露"),
 ("desc_cn","中文一句话定位","本轮重写"),
 ("tag_l1/tag_l2","一/二级标签(V20)","另一AI维护，只读"),
 ("business_tags.role","企业角色","本轮补全"),
 ("highlights","亮点列表","本轮重写差异化事实"),
 ("events","重大事件时间线","本轮新增"),
 ("payor_model","付费方模式","本轮补全"),
 ("silver_verdict","银发相关性","本轮判定 核心银发/泛医疗擦边/非银发"),
 ("signal_strength","信号强度0~10","已算好(脚本)"),
 ("info_score/diff_score/copy_score","四维评分0~10","本轮强模评"),
 ("research_value","综合分0~100","=(信号×.3+信息×.3+差异×.2+复制×.2)×10"),
 ("recommend","三版推荐理由","本轮重做"),
 ("update_time","更新时间","2026-07-17"),
]
ws5.append(["字段","含义","处置说明"])
for row in FIELD_DICT: ws5.append(list(row))
for c in range(1,4):
    ws5.cell(1,c).fill=hdr_fill; ws5.cell(1,c).font=hdr_font; ws5.cell(1,c).alignment=ctr
ws5.column_dimensions["A"].width=34; ws5.column_dimensions["B"].width=52; ws5.column_dimensions["C"].width=40
for r in range(2, ws5.max_row+1):
    ws5.cell(r,2).alignment=wrap; ws5.cell(r,3).alignment=wrap

OUT = os.path.join(BASE, "scores_draft", f"SilverPulse_评分补全_全量_{len(proc)}家_2026-07-17.xlsx")
wb.save(OUT)
print("Excel 已生成:", OUT)
print("Sheet1 行:", ws1.max_row, " Sheet2 行:", ws2.max_row, " Sheet3 行:", ws3.max_row, " Sheet4 行:", ws4.max_row)
print("处理企业数:", len(proc), " 标签建议:", len(tag_review), " 非银发:", len(nonsilver))
