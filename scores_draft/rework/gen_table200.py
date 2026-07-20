# -*- coding: utf-8 -*-
"""把已返工的 200 家企业全部字段导出为 Excel，并逐家跑机器门禁嵌入「验证状态」列。"""
import json, glob, re, sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import check_single as C
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
BD = os.path.join(os.path.dirname(os.path.abspath(__file__)), "batches200")
DB = os.path.join(ROOT, "data/enterprise/all_enterprises.json")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "返工企业全字段_200家.xlsx")

# ---- 1. 取 200 家序列号 ----
serials = []
for f in sorted(glob.glob(BD + "/batch_src_*.json")):
    data = json.load(open(f, encoding="utf-8"))
    for e in data:
        serials.append(e.get("serial"))
db = {e["serial"]: e for e in json.load(open(DB, encoding="utf-8"))}
recs = [db[s] for s in serials]
assert len(recs) == 200, f"期望200家，实得{len(recs)}"

# ---- 2. 复校（R1-R9 用门禁；R10 跨企业用 trigram 重合率，快速） ----
def trigrams(s):
    s = re.sub(r"\s", "", s or "")
    return set(s[i:i+3] for i in range(len(s) - 2))

tg = {e["serial"]: trigrams(e.get("recommend", "")) for e in recs}

def cross_dup(e):
    t = trigrams(e.get("recommend", ""))
    if not t:
        return None
    for o in recs:
        if o is e:
            continue
        so = tg.get(o["serial"])
        if so and len(t & so) / len(t | so) > 0.6:
            return o["serial"]
    return None

verify = {}  # serial -> (status, issues_list)
for e in recs:
    errs = C.validate(e, others=None, skip={"R10"})
    cd = cross_dup(e)
    if cd:
        errs.append(f"R10:跨企业雷同(与{cd}重合>0.6)")
    verify[e["serial"]] = ("FAIL" if errs else "PASS", errs)

# ---- 3. 组装表格 ----
flat_cols = [
    ("serial", "序列号"), ("name", "英文名"), ("name_cn", "中文名"),
    ("tag_l1", "一级标签"), ("tag_l2", "二级标签"), ("region", "地区"),
    ("category_l1", "分类一级"), ("category_l2", "分类二级"),
    ("business_model", "商业模式"), ("payor_model", "付费方"),
    ("silver_verdict", "银发判定"), ("source", "来源"), ("数据来源", "数据来源"),
    ("update_time", "更新时间"),
    ("signal_strength", "信号强度"), ("info_score", "信息量"), ("diff_score", "差异化"),
    ("copy_score", "可复制"), ("research_value", "综合分"), ("value_score", "价值分"),
    ("desc_cn", "企业描述"), ("silver_reason", "银发理由"), ("recommend", "推荐理由(单字符串)"),
    ("验证状态", "验证状态"), ("问题明细", "问题明细"),
]
nested_keys = ["business_tags", "funding_latest", "funding_total", "events",
               "news_coverage", "tags", "highlights", "investors", "stage",
               "founded", "website_url", "crunchbase_url", "description"]

wb = Workbook()
ws = wb.active
ws.title = "企业全字段"
hdr_font = Font(bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="305496")
for j, (key, label) in enumerate(flat_cols, 1):
    c = ws.cell(1, j, label)
    c.font = hdr_font
    c.fill = hdr_fill
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
for i, e in enumerate(recs, 2):
    for j, (key, label) in enumerate(flat_cols, 1):
        if key == "验证状态":
            v = verify[e["serial"]][0]
        elif key == "问题明细":
            v = "；".join(verify[e["serial"]][1])
        else:
            v = e.get(key)
            if isinstance(v, list):
                v = "、".join(str(x) for x in v)
            elif v is None:
                v = ""
        c = ws.cell(i, j, v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if key == "验证状态":
            c.font = Font(bold=True, color=("C00000" if v == "FAIL" else "375623"))
# 列宽
widths = {"serial":10,"name":16,"name_cn":16,"tag_l1":12,"tag_l2":16,"region":8,
          "category_l1":12,"category_l2":14,"business_model":20,"payor_model":18,
          "silver_verdict":12,"source":16,"数据来源":12,"update_time":12,
          "signal_strength":9,"info_score":8,"diff_score":8,"copy_score":8,
          "research_value":9,"value_score":8,"desc_cn":50,"silver_reason":40,
          "recommend":60,"验证状态":10,"问题明细":30}
for j,(key,label) in enumerate(flat_cols,1):
    ws.column_dimensions[chr(64+j) if j<=26 else "A"+chr(64+j-26)].width = widths.get(key,14)
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{chr(64+len(flat_cols)) if len(flat_cols)<=26 else 'A'+chr(64+len(flat_cols)-26)}201"

# 嵌套字段 sheet
ws2 = wb.create_sheet("嵌套字段JSON")
ws2.cell(1,1,"序列号"); ws2.cell(1,2,"英文名")
for j,k in enumerate(nested_keys,3):
    ws2.cell(1,j,k)
for i,e in enumerate(recs,2):
    ws2.cell(i,1,e.get("serial"))
    ws2.cell(i,2,e.get("name",e.get("name_cn")))
    for j,k in enumerate(nested_keys,3):
        v = e.get(k)
        ws2.cell(i,j, json.dumps(v, ensure_ascii=False) if not isinstance(v,str) else v)
ws2.freeze_panes = "A2"

# 总览 sheet
ws3 = wb.create_sheet("验证总览", 0)
n_pass = sum(1 for s,_ in verify.values() if s == "PASS")
fails = [(s, errs) for s,(st,errs) in verify.items() if st == "FAIL"]
rvs = [e.get("research_value",0) for e in recs]
wc = [C.content_len(e.get("recommend","")) for e in recs]
summary = [
    ("已返工企业总数", 200),
    ("机器门禁 PASS", n_pass),
    ("机器门禁 FAIL", len(fails)),
    ("综合分(research_value)均值", round(sum(rvs)/len(rvs),1)),
    ("综合分最低", min(rvs)),
    ("综合分最高", max(rvs)),
    ("推荐理由字数均值", round(sum(wc)/len(wc),1)),
    ("推荐理由字数最短", min(wc)),
    ("推荐理由字数最长", max(wc)),
    ("门禁阈值", "R2 70~220字 / R1 单字符串 / R5 四维 / R-name 禁名 / R-dedup 禁抄desc / R-jargon 禁黑话 / R-noabs 禁绝对化 / R6 desc≥80 / R7 silver≥30 / R8 payor规范 / R10 跨企业重合≤0.6"),
]
ws3.cell(1,1,"返工企业200家 · 自检总览").font = Font(bold=True, size=14)
for i,(k,v) in enumerate(summary,3):
    ws3.cell(i,1,k).font = Font(bold=True)
    ws3.cell(i,2,v)
ws3.cell(len(summary)+4,1,"FAIL 明细：").font = Font(bold=True, color="C00000")
r = len(summary)+5
for s,errs in fails:
    ws3.cell(r,1,s)
    ws3.cell(r,2,"；".join(errs))
    r += 1
ws3.column_dimensions["A"].width = 30
ws3.column_dimensions["B"].width = 90

wb.save(OUT)
print("SAVED:", OUT)
print(f"PASS={n_pass} FAIL={len(fails)}")
if fails:
    for s,errs in fails:
        print("  FAIL",s,errs)
