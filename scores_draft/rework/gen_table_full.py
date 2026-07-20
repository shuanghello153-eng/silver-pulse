# -*- coding: utf-8 -*-
"""
全量1502家企业 · 全字段Excel生成器 V4版。
对照《企业字段字典_V1.md》补全所有关键字段为扁平列，嵌套JSON作辅助sheet。
含V4门禁验证状态列。用法：python gen_table_full.py
"""
import json, os, sys, re, glob
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import check_single as C
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.normpath(os.path.join(HERE, "..", ".."))
DB = os.path.join(BASE, "data/enterprise/all_enterprises.json")
OUT = os.path.join(HERE, "全量企业_1502家_V4.xlsx")

# ── 扁平列定义（按字典V1分组，优先级排序）──
FLAT_COLS = [
    # A. 基础信息
    ("serial", "序列号"), ("name", "英文名"), ("name_cn", "中文名"),
    ("region", "地区"), ("founded", "成立年份"), ("stage", "阶段"),
    ("website_url", "官网"),
    # B. 投融资
    ("investors", "投资方"),
    ("funding_round", "最新轮次"), ("funding_amount", "最新金额"),
    ("funding_date", "最新时间"), ("funding_display", "融资展示串"),
    ("funding_total_display", "累计融资展示"),
    # C. 标签与分类
    ("tag_l1", "一级标签"), ("tag_l2", "二级标签"),
    ("category_l1", "旧分类一级"), ("category_l2", "旧分类二级"),
    ("silver_verdict", "银发判定"),
    # D. 商业模式
    ("business_model", "商业模式"), ("payor_model", "付费方"),
    ("source", "来源"), ("数据来源", "数据来源"), ("update_time", "更新时间"),
    # E. 评分（四维+综合）
    ("signal_strength", "信号强度"), ("info_score", "信息量"),
    ("diff_score", "差异化"), ("copy_score", "可复制"),
    ("research_value", "综合分"), ("value_score", "价值分(旧)"),
    # F. 文本字段
    ("desc_cn", "企业描述"), ("silver_reason", "银发理由"),
    ("recommend", "推荐理由(单字符串)"),
    # G. 验证状态
    ("验证状态", "验证状态"), ("问题明细", "问题明细"),
]

NESTED_KEYS = ["business_tags", "business_model_cn", "highlights", "events",
               "news_coverage", "tags", "crunchbase_url", "description",
               "funding_latest", "funding_total"]


def _flat_val(e, key):
    """从企业记录中提取扁平值。"""
    if key == "funding_round":
        fl = e.get("funding_latest") or {}
        return fl.get("round", "") if isinstance(fl, dict) else ""
    if key == "funding_amount":
        fl = e.get("funding_latest") or {}
        return fl.get("amount", "") if isinstance(fl, dict) else ""
    if key == "funding_date":
        fl = e.get("funding_latest") or {}
        return fl.get("date", "") if isinstance(fl, dict) else ""
    if key == "funding_display":
        fl = e.get("funding_latest") or {}
        return fl.get("display", "") if isinstance(fl, dict) else ""
    if key == "funding_total_display":
        ft = e.get("funding_total") or {}
        return ft.get("display", "") if isinstance(ft, dict) else ""
    v = e.get(key)
    if isinstance(v, list):
        return "、".join(str(x) for x in v)
    return "" if v is None else str(v)


def main():
    db = json.load(open(DB, encoding="utf-8"))
    db_sorted = sorted(db, key=lambda x: x.get("research_value") or -1, reverse=True)
    print(f"全库 {len(db_sorted)} 家，按 research_value 降序")

    # ── V4 门禁复校（单企业内部，不含R10）──
    print("运行V4门禁...")
    verify = {}
    for i, e in enumerate(db_sorted):
        errs = C.validate(e, others=None, skip={"R10"})
        status = "FAIL" if errs else "PASS"
        detail = "；".join(errs) if errs else ""
        verify[e["serial"]] = (status, detail)
        if (i + 1) % 200 == 0:
            print(f"  已校 {i+1}/{len(db_sorted)}")

    n_pass = sum(1 for s, _ in verify.values() if s == "PASS")
    print(f"门禁结果: PASS={n_pass} FAIL={len(db_sorted)-n_pass}")

    # ── 写Excel ──
    wb = Workbook()

    # Sheet 1: 验证总览
    ws0 = wb.active
    ws0.title = "验证总览"
    ws0.cell(1, 1, "银发经济企业库 V4 全量验证总览").font = Font(bold=True, size=14)
    summary = [
        ("全库总数", len(db_sorted)), ("V4-PASS", n_pass), ("V4-FAIL", len(db_sorted) - n_pass),
        ("门禁版本", "V4 (2026-07-18)"),
        ("新规则", "R-field-dedup(6字段去重≤10字) + R-integrity(内部一致性) + R-novelty(增量≥30%) + R10跨企业(重合≤0.5)"),
    ]
    rvs = [e.get("research_value") or 0 for e in db_sorted]
    rcl = [C.content_len(e.get("recommend", "")) for e in db_sorted if isinstance(e.get("recommend"), str)]
    summary.extend([
        ("综合分均值", round(sum(rvs)/len(rvs),1)), ("综合分最低", min(rvs)), ("综合分最高", max(rvs)),
        ("推荐理由字数均值(仅str)", round(sum(rcl)/len(rcl),1) if rcl else 0),
    ])
    for i, (k, v) in enumerate(summary, 3):
        ws0.cell(i, 1, k).font = Font(bold=True); ws0.cell(i, 2, v)
    # FAIL明细
    r0 = len(summary) + 5
    ws0.cell(r0, 1, "FAIL 企业清单").font = Font(bold=True, color="C00000")
    r0 += 1
    for s in sorted(db_sorted):
        st, dt = verify[s["serial"]]
        if st == "FAIL":
            ws0.cell(r0, 1, s["serial"]); ws0.cell(r0, 2, dt); r0 += 1
    ws0.column_dimensions["A"].width = 30; ws0.column_dimensions["B"].width = 120

    # Sheet 2: 企业全字段
    ws = wb.create_sheet("企业全字段")
    hdr_font = Font(bold=True, color="FFFFFF"); hdr_fill = PatternFill("solid", fgColor="305496")
    for j, (key, label) in enumerate(FLAT_COLS, 1):
        c = ws.cell(1, j, label); c.font = hdr_font; c.fill = hdr_fill
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for i, e in enumerate(db_sorted, 2):
        for j, (key, _) in enumerate(FLAT_COLS, 1):
            if key == "验证状态":
                v = verify[e["serial"]][0]
            elif key == "问题明细":
                v = verify[e["serial"]][1]
            else:
                v = _flat_val(e, key)
            c = ws.cell(i, j, v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if key == "验证状态" and v == "FAIL":
                c.font = Font(bold=True, color="C00000")
            elif key == "验证状态" and v == "PASS":
                c.font = Font(bold=True, color="375623")
    widths = {"serial":10,"name":16,"name_cn":16,"region":8,"founded":8,"stage":10,
              "website_url":22,"investors":16,"funding_round":12,"funding_amount":16,
              "funding_date":10,"funding_display":30,"funding_total_display":20,
              "tag_l1":12,"tag_l2":18,"category_l1":12,"category_l2":14,
              "silver_verdict":12,"business_model":24,"payor_model":18,
              "source":16,"数据来源":12,"update_time":12,
              "signal_strength":9,"info_score":8,"diff_score":8,"copy_score":8,
              "research_value":9,"value_score":9,
              "desc_cn":50,"silver_reason":40,"recommend":65,
              "验证状态":10,"问题明细":35}
    last_col = chr(64+len(FLAT_COLS)) if len(FLAT_COLS)<=26 else 'A'+chr(64+len(FLAT_COLS)-26)
    for j,(key,_) in enumerate(FLAT_COLS,1): 
        cid=chr(64+j) if j<=26 else f'A{chr(64+j-26)}'
        ws.column_dimensions[cid].width = widths.get(key,14)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:{last_col}{len(db_sorted)+1}"

    # Sheet 3: 嵌套字段
    ws2 = wb.create_sheet("嵌套字段JSON")
    ws2.cell(1,1,"序列号"); ws2.cell(1,2,"英文名")
    for j,k in enumerate(NESTED_KEYS,3): ws2.cell(1,j,k)
    for i,e in enumerate(db_sorted,2):
        ws2.cell(i,1,e["serial"]); ws2.cell(i,2,e.get("name",""))
        for j,k in enumerate(NESTED_KEYS,3):
            v=e.get(k); ws2.cell(i,j,json.dumps(v,ensure_ascii=False) if not isinstance(v,str) else v)

    wb.save(OUT)
    print(f"\nSAVED: {OUT}")
    print(f"大小: {os.path.getsize(OUT)/1024/1024:.1f} MB")


if __name__ == "__main__":
    main()
