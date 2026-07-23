# -*- coding: utf-8 -*-
import json, os, collections
ROOT="G:/workbuddy/2026-06-28-23-34-20/silver-pulse"
DB=os.path.join(ROOT,"data/enterprise/all_enterprises.json")
BEFORE=os.path.join(ROOT,"scores_draft/pilot25_before.json")
OUT=os.path.join(ROOT,"scores_draft/Pilot25_交付_改进与字段总表.xlsx")

data=json.load(open(DB,encoding="utf-8"))
by={e["serial"]:e for e in data}
before=json.load(open(BEFORE,encoding="utf-8"))
serials=json.load(open(os.path.join(ROOT,"scores_draft/pilot25_serials.json"),encoding="utf-8"))

# ---------- 1) 全库字段填充率 ----------
n=len(data)
present=collections.Counter(); nonempty=collections.Counter()
for e in data:
    for k,v in e.items():
        present[k]+=1
        if v not in (None,"",[],{},0) and not (isinstance(v,str) and v.strip()==""):
            nonempty[k]+=1
all_fields=sorted(present)

# 字段含义 / 处置 / 理由（本项目视角）
META={
 "serial":("唯一序号","系统字段，入库即定，不动","—"),
 "name":("企业英文名","系统字段，不动（打标AI录入）","—"),
 "name_cn":("企业中文名","系统字段，不动（打标AI维护）","—"),
 "region":("地区（海外/国内）","不动（打标/采集维护）","—"),
 "founded":("成立年份","✅ 本项目补全：联网核实，纠正错填","评分用；错填会误导年龄判断"),
 "stage":("发展阶段（初创/成长/上市等）","不动（打标维护）","—"),
 "source":("数据来源标注（哪个信源/AI录入）","不动（记录录入出处）","—"),
 "website_url":("官网链接","不动","—"),
 "desc_cn":("中文一句话简介","✅ 重写：用联网事实填实质业务，替换模板废话","卡片+推荐直接展示"),
 "description":("英文简介","低优先：不打标、非空则不碰","与desc_cn重复，价值低"),
 "funding_latest":("最新一轮融资/收购","✅ 重点补全：纠正 amount 被错填成投资方名 的老问题，补齐轮次/金额/时间/投资方","选题热点核心信号"),
 "funding_total":("累计融资额","✅ 补全（搜得到才填）","判断体量"),
 "investors":("投资方列表","✅ 补全","判断背书质量"),
 "payor_model":("支付方（谁出钱）","✅ 新增维度：政府/商保/个人/C端/混合，附简述","国内借鉴最关键差异点"),
 "tag_l1":("一级标签","🔴 只读不写（另一AI专职维护）","避免与打标工作冲突"),
 "tag_l2":("二级标签","🔴 只读不写（另一AI专职维护）","避免与打标工作冲突"),
 "business_tags":("结构化标签（含 role 等）","只读：role 作输入，不主动补；高把握才标【动你标注】","V20已替代其职能"),
 "business_model":("商业模式自由文本","🔴 废弃：不主动重写1210条。重复度高、辨识度低、V20标签已替代、推荐理由不读它","重写=黑洞成本，零产出"),
 "business_model_cn":("商业模式中文","🔴 废弃：同上","同上"),
 "category_l1":("旧一级分类","🔴 废弃/不维护：前端已不读，与tag体系重复","保留兼容，零风险派生"),
 "category_l2":("旧二级分类","🔴 废弃/不维护：前端已不读","同上"),
 "news_coverage":("新闻覆盖情况","不动（另一流程负责）；本研究发现的缺口只提示，不主动写","避免越界"),
 "highlights":("亮点标签","✅ 重写：基于事实提炼2-3条","卡片展示"),
 "value_score":("旧综合评分","✅ 桥接写入 = research_value（旧前端回退兼容）","保证旧UI不崩"),
 "research_value":("综合研究价值分 0-100","✅ 本项目新增：四阶段加权（信号×0.3+信息×0.3+差异×0.2+复制×0.2）×10","推荐排序主依据"),
 "signal_strength":("信号强度 0-10","✅ 新增（阶段3规则脚本算）","初筛"),
 "info_score":("信息量 0-10","✅ 新增（阶段4·强模判定）","研究素材充足度"),
 "diff_score":("差异化/反共识 0-10","✅ 新增（阶段4·强模判定）","选题故事性"),
 "copy_score":("可复制/国内借鉴 0-10","✅ 新增（阶段4·强模判定）","创业者价值落点"),
 "recommend":("推荐理由（详细）","✅ 新增：结合四维+公开事实，不重复卡片已有信息","小爽选题决策直接依据"),
 "update_time":("信息更新时间","✅ 新增：ISO日期，记录本轮补全","溯源"),
 "ingest_time":("录入时间（另一AI加）","不动、不覆盖","他人字段"),
 "needs_review":("待复核标记","不动","他人字段"),
}
def meta(k):
    return META.get(k,("（待补充含义）","（待定）","（待定）"))

def trunc(s,lim=60):
    if s is None: return ""
    s=str(s).replace("\n"," ")
    return s if len(s)<=lim else s[:lim]+"…"

# ---------- 2) 改进对比（25家） ----------
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

wb=Workbook()
hdr_fill=PatternFill("solid",fgColor="1F4E78")
hdr_font=Font(color="FFFFFF",bold=True,size=10)
sub_fill=PatternFill("solid",fgColor="DDEBF7")
red_fill=PatternFill("solid",fgColor="FCE4D6")
wrap=Alignment(wrap_text=True,vertical="top")
thin=Side(style="thin",color="BFBFBF")
border=Border(left=thin,right=thin,top=thin,bottom=thin)

def style_header(ws,row,ncol):
    for c in range(1,ncol+1):
        cell=ws.cell(row=row,column=c)
        cell.fill=hdr_fill; cell.font=hdr_font; cell.alignment=wrap; cell.border=border

# Sheet1 改进情况对比
ws1=wb.active; ws1.title="改进情况对比"
cols1=["序号","企业(中)","融资_改前","融资_改后","支付方_改前","支付方_改后",
        "描述_改前(截断)","描述_改后(截断)","评分_改前","评分_改后","更新时间_改前","更新时间_改后"]
ws1.append(cols1); style_header(ws1,1,len(cols1))
for s in serials:
    e=by[s]; b=before.get(s,{})
    fl=e.get("funding_latest") or {}
    flb=b.get("funding_latest") or {}
    pm=e.get("payor_model") or ""
    pmb=b.get("payor_model") or ""
    rv=e.get("research_value"); rvb=b.get("research_value")
    ut=e.get("update_time") or ""; utb=b.get("update_time") or ""
    row=[s, e.get("name_cn") or e.get("name"),
         trunc(flb.get("display") if isinstance(flb,dict) else flb,40),
         trunc(fl.get("display") if isinstance(fl,dict) else fl,40),
         trunc(pmb,24), trunc(pm,24),
         trunc(b.get("desc_cn"),50), trunc(e.get("desc_cn"),50),
         rvb if rvb is not None else "无",
         rv if rv is not None else "无",
         trunc(utb,12), trunc(ut,12)]
    ws1.append(row)
widths1=[8,16,30,30,18,18,42,42,9,9,12,12]
for i,w in enumerate(widths1,1): ws1.column_dimensions[get_column_letter(i)].width=w
for r in range(2,ws1.max_row+1):
    for c in range(1,len(cols1)+1):
        cell=ws1.cell(row=r,column=c); cell.alignment=wrap; cell.border=border
ws1.freeze_panes="A2"

# Sheet2 评分与推荐理由（全文本）
ws2=wb.create_sheet("评分与推荐理由")
cols2=["序号","企业(中)","信号","信息","差异","复制","综合分","推荐理由（详细）"]
ws2.append(cols2); style_header(ws2,1,len(cols2))
for s in serials:
    e=by[s]
    row=[s,e.get("name_cn") or e.get("name"),
         e.get("signal_strength"),e.get("info_score"),e.get("diff_score"),e.get("copy_score"),
         e.get("research_value"), e.get("recommend") or ""]
    ws2.append(row)
widths2=[8,16,7,7,7,7,9,90]
for i,w in enumerate(widths2,1): ws2.column_dimensions[get_column_letter(i)].width=w
for r in range(2,ws2.max_row+1):
    for c in range(1,len(cols2)+1):
        cell=ws2.cell(row=r,column=c); cell.alignment=wrap; cell.border=border
ws2.freeze_panes="A2"

# Sheet3 企业库字段全表
ws3=wb.create_sheet("企业库字段全表")
cols3=["字段名","中文含义","全库填充率","处置","理由"]
ws3.append(cols3); style_header(ws3,1,len(cols3))
for k in all_fields:
    mean,disp,reason=meta(k)
    rate=f"{nonempty[k]}/{n} ({round(100*nonempty[k]/n)}%)"
    ws3.append([k,mean,rate,disp,reason])
    if "🔴" in disp:
        for c in range(1,len(cols3)+1):
            ws3.cell(row=ws3.max_row,column=c).fill=red_fill
widths3=[20,34,18,40,46]
for i,w in enumerate(widths3,1): ws3.column_dimensions[get_column_letter(i)].width=w
for r in range(2,ws3.max_row+1):
    for c in range(1,len(cols3)+1):
        cell=ws3.cell(row=r,column=c); cell.alignment=wrap; cell.border=border
ws3.freeze_panes="A2"

wb.save(OUT)
print("SAVED",OUT)
print("fields total:",len(all_fields))
print("rows in 改进情况:",ws1.max_row-1)
print("rows in 评分理由:",ws2.max_row-1)
print("rows in 字段全表:",ws3.max_row-1)
