# -*- coding: utf-8 -*-
"""生成三表 Excel：Sheet1 重要字段修改前后对比；Sheet2 更新企业全部真实字段；Sheet3 字段字典。"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = "G:/workbuddy/2026-06-28-23-34-20/silver-pulse"
SD = os.path.join(BASE, "scores_draft")

def load(p):
    with open(p, encoding="utf-8") as f:
        return json.load(f)

new30 = load(os.path.join(SD, "new30_serials.json"))
pilot = load(os.path.join(SD, "pilot25_serials.json"))
pilot_before = load(os.path.join(SD, "pilot25_before.json"))
before_snap = load(os.path.join(SD, "before_snapshot.json"))
after_snap = load(os.path.join(SD, "after_snapshot.json"))

all_serials = new30 + pilot

def jdump(v):
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    return "" if v is None else str(v)

def name_of(e):
    n = e.get("name", "") or ""
    nc = e.get("name_cn", "") or ""
    if nc and nc != n:
        return f"{n}（{nc}）" if n else nc
    return n

# ============ Sheet1 ============
wb = Workbook()
ws1 = wb.active
ws1.title = "Sheet1_前后对比"
headers1 = ["编号", "企业名称", "一级标签", "银发判定",
            "信号强度(前)", "信号强度(后)", "研究价值(前)", "研究价值(后)",
            "信息量", "差异化", "可复制",
            "一句话定位(后)", "推荐理由 v1(后)", "推荐理由 v2(后)", "推荐理由 v3(后)"]
ws1.append(headers1)

for s in all_serials:
    a = after_snap.get(s, {})
    b = before_snap.get(s, {}) if s in new30 else pilot_before.get(s, {})
    # 前：30-new 用 before_snap；pilot 用 pilot25_before
    b_src = before_snap.get(s, {}) if s in new30 else pilot_before.get(s, {})
    rec_a = a.get("recommend", {}) or {}
    if isinstance(rec_a, str):
        rec_a = {}
    silver = a.get("silver_verdict") or b_src.get("silver_verdict") or ""
    row = [
        s,
        name_of(a),
        "、".join(a.get("tag_l1") or []) or "、".join(b_src.get("tag_l1") or []),
        silver,
        jdump(b_src.get("signal_strength")),
        jdump(a.get("signal_strength")),
        jdump(b_src.get("research_value")),
        jdump(a.get("research_value")),
        jdump(rec_a.get("info_score")),
        jdump(rec_a.get("diff_score")),
        jdump(rec_a.get("copy_score")),
        jdump(a.get("desc_cn")),
        jdump(rec_a.get("rec_v1")),
        jdump(rec_a.get("rec_v2")),
        jdump(rec_a.get("rec_v3")),
    ]
    ws1.append(row)

# 样式
hdr_fill = PatternFill("solid", fgColor="1F4E78")
hdr_font = Font(bold=True, color="FFFFFF", size=10)
green_fill = PatternFill("solid", fgColor="E2EFDA")
for c in range(1, len(headers1) + 1):
    cell = ws1.cell(row=1, column=c)
    cell.fill = hdr_fill; cell.font = hdr_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
# 推荐理由列宽 + 换行
widths1 = [9, 26, 14, 11, 11, 11, 11, 11, 7, 7, 7, 40, 46, 46, 46]
for i, w in enumerate(widths1, 1):
    ws1.column_dimensions[get_column_letter(i)].width = w
for r in range(2, ws1.max_row + 1):
    for c in [12, 13, 14, 15]:
        ws1.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    # 后值标绿
    for c in [6, 8]:
        ws1.cell(row=r, column=c).fill = green_fill
ws1.freeze_panes = "A2"

# ============ Sheet2 ============
ws2 = wb.create_sheet("Sheet2_全字段")
# 收集所有键，按逻辑排序
preferred = ["serial", "name", "name_cn", "region", "founded", "hq", "stage",
             "tag_l1", "tag_l2", "business_tags", "category_l1", "category_l2",
             "description", "desc_cn", "highlights", "events",
             "funding_latest", "funding_total", "investors", "payor_model",
             "news_coverage", "source", "website_url", "crunchbase_url",
             "business_model", "business_model_cn", "tags",
             "signal_strength", "research_value", "recommend",
             "silver_verdict", "silver_reason", "update_time", "value_score"]
allkeys = set()
for s in all_serials:
    allkeys.update(after_snap.get(s, {}).keys())
extra = [k for k in allkeys if k not in preferred]
cols2 = preferred + sorted(extra)
ws2.append(cols2)
for s in all_serials:
    e = after_snap.get(s, {})
    ws2.append([jdump(e.get(k)) for k in cols2])
for c in range(1, len(cols2) + 1):
    cell = ws2.cell(row=1, column=c)
    cell.fill = hdr_fill; cell.font = hdr_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws2.column_dimensions[get_column_letter(c)].width = 22
# 长文本列宽
for ci, k in enumerate(cols2, 1):
    if k in ("desc_cn", "description", "highlights", "events", "recommend", "news_coverage", "business_tags"):
        ws2.column_dimensions[get_column_letter(ci)].width = 50
ws2.freeze_panes = "B2"

# ============ Sheet3 字段字典 ============
ws3 = wb.create_sheet("Sheet3_字段字典")
FIELD_DICT = [
    ("serial", "企业内部编号 #0001 起", "主键，不动"),
    ("name / name_cn", "英文名 / 中文名", "name 不动；name_cn 可补中文译名"),
    ("region", "地区：国内 / 海外", "不动，影响可复制判断"),
    ("founded", "成立年份", "本轮联网核实纠正"),
    ("stage", "成长阶段", "按实际填：已上市/被收购/融资中/未披露"),
    ("description / desc_cn", "英文 / 中文简介", "本轮重写 desc_cn（一句话定位）"),
    ("funding_latest", "最新一轮融资/收购 {round,amount,currency,date,display,investors}", "结构纠正+金额核实"),
    ("funding_total", "累计融资额", "补真实累计；上市公司写不适用"),
    ("investors", "最新一轮投资方[]", "只记最新一轮"),
    ("business_tags", "{customer,role,channel}", "role 本轮据研究微调"),
    ("tag_l1 / tag_l2", "一级/二级标签(V20)", "另一 AI 维护，只读"),
    ("highlights", "亮点列表", "本轮重写为有信息量的差异化事实"),
    ("events", "重大事件时间线[]", "本轮新增，补 IPO/收购/合作等"),
    ("payor_model", "付费方模式", "本轮补：政府/商保/个人/B端/混合"),
    ("news_coverage", "{news_count,news_quality,latest_news[]}", "保留爬虫数据；web 新闻仅追加不覆盖"),
    ("silver_verdict", "银发相关性：核心银发/泛医疗擦边/非银发", "本轮新增判定"),
    ("signal_strength", "阶段3 信号强度 0~10（脚本）", "本轮新增；IPO+6、融资/收购额分级、时效"),
    ("info_score", "阶段4 信息量 0~10", "本轮新增（强模评）"),
    ("diff_score", "阶段4 差异化 0~10", "本轮新增（强模评）"),
    ("copy_score", "阶段4 可复制 0~10", "本轮新增（强模评）"),
    ("research_value", "综合评分 0~100", "=(信号×30%+信息量×30%+差异化×20%+可复制×20%)×10"),
    ("recommend", "{rec_v1,rec_v2,rec_v3,info/diff/copy_score,signal,rv}", "本轮重做为三版本推荐理由"),
    ("update_time", "更新时间 ISO 日期", "本轮新增，如 2026-07-16"),
    ("category_l1/l2 / tags / business_model", "旧分类/旧标签/旧商业模式", "已废弃，不主动维护"),
]
ws3.append(["字段", "含义", "处置说明"])
for row in FIELD_DICT:
    ws3.append(list(row))
for c in range(1, 4):
    cell = ws3.cell(row=1, column=c)
    cell.fill = hdr_fill; cell.font = hdr_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws3.column_dimensions["A"].width = 34
ws3.column_dimensions["B"].width = 52
ws3.column_dimensions["C"].width = 40
for r in range(2, ws3.max_row + 1):
    ws3.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws3.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="top")

OUT = os.path.join(BASE, "scores_draft", "SilverPulse_评分信息补全_55家_2026-07-16.xlsx")
wb.save(OUT)
print("Excel 已生成：", OUT)
print("Sheet1 行数(含表头):", ws1.max_row, " Sheet2 列数:", len(cols2), " Sheet2 行数:", ws2.max_row)
