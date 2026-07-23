# -*- coding: utf-8 -*-
# 复盘证据采集：字段完整性 / 表格遗漏 / 推荐重复 / Medicare长护错杀
import json, re
from collections import Counter

DB = json.load(open('../../data/enterprise/all_enterprises.json', encoding='utf-8'))
proc = set(json.load(open('processed.json', encoding='utf-8')))
N = len(DB)

# 1) 全字段 schema + 空值率
allkeys = Counter()
for e in DB:
    for k in e.keys():
        allkeys[k] += 1
print("=== 1) DB 全字段（%d 家企业，出现次数）===" % N)
empty = {}
for k in allkeys:
    em = sum(1 for e in DB if e.get(k) in (None, '', [], {}))
    empty[k] = em
for k, c in allkeys.most_common():
    print(f"  {k:22s} 出现{c:4d}  空值{empty[k]:4d} ({round(empty[k]/N*100,1)}%)")

# 2) 我发的增强Excel到底有哪些列（直接读xlsx）
try:
    from openpyxl import load_workbook
    wb = load_workbook('G:/workbuddy/2026-07-17-12-07-23/SilverPulse_增强版_1502家_字段完整版.xlsx')
    ws = wb['全量总览']
    cols = [c.value for c in ws[1]]
    print("\n=== 2) 增强Excel「全量总览」实际列 ===")
    print("  ", cols)
    missing = [k for k in ['desc_cn','tag_l1','tag_l2','business_tags','founded','website_url','funding_latest','funding_total','investors','highlights','events','business_tags_role','silver_reason'] if k not in cols]
    print("  >> 用户关心但表格遗漏的字段:", missing)
except Exception as ex:
    print("读xlsx失败:", ex)

# 3) 推荐理由重复率
recs = [e.get('recommend') or '' for e in DB]
uniq = set(recs)
print("\n=== 3) recommend 重复率 ===")
print(f"  总 {len(recs)}  唯一 {len(uniq)}  重复率 {round((1-len(uniq)/len(recs))*100,1)}%")
# “切入X赛道”分布
cut = Counter()
for r in recs:
    for x in re.findall(r'切入([^（(，。；]+)', r):
        cut['切入'+x.strip()] += 1
print("  >> “切入X赛道” top15:")
for t, c in cut.most_common(15):
    print(f"     {c:4d}  {t}")
# 收尾括号内容
tail = Counter()
for r in recs:
    m = re.search(r'[（(]([^）)]*)[）)]$', r)
    if m:
        tail['(末括号)'+m.group(1)[:24]] += 1
    else:
        tail['末10字:'+r[-10:]] += 1
print("  >> 收尾模式 top15:")
for t, c in tail.most_common(15):
    print(f"     {c:4d}  {t}")
# 含“结合国内/国内资源/本地资源”等套话
cliche = sum(1 for r in recs if re.search(r'结合国内|国内资源|本地资源|可借鉴落地|可直接借鉴', r))
print(f"  >> 含“结合国内/本地资源/可借鉴落地”等套话: {cliche} ({round(cliche/N*100,1)}%)")

# 4) Medicare / 长护 / 老年医保 错杀检查
print("\n=== 4) Medicare/长护/老年医保 相关检索 ===")
hits = []
for e in DB:
    blob = ' '.join(str(e.get(k) or '') for k in ['name','name_cn','desc_cn','recommend','tag_l1','tag_l2','business_tags','silver_reason','highlights'])
    if re.search(r'Medicare|医保|长护|长期护理|long.?term care|老年医保|CMS|senior care', blob, re.I):
        hits.append((e['serial'], e.get('name'), e.get('silver_verdict'), e.get('tag_l1'), e.get('tag_l2')))
print(f"  命中 {len(hits)} 家")
noncore = [h for h in hits if h[2] in ('非银发','泛医疗擦边')]
print(f"  其中被标 非银发/泛医疗擦边 的: {len(noncore)} 家  ⚠️ 需人工复核是否错杀")
for h in noncore[:40]:
    print(f"     {h[0]} {h[1]} | {h[2]} | L1={h[3]} L2={h[4]}")

# 5) 增强 vs 基线 + 银发分布
print("\n=== 5) 覆盖与银发判定 ===")
print("  增强(已联网精评):", sum(1 for e in DB if e['serial'] in proc))
print("  基线(未联网):", sum(1 for e in DB if e['serial'] not in proc))
sv = Counter(e.get('silver_verdict') or '（未标注）' for e in DB)
print("  银发判定:", dict(sv))

# 6) 抽样增强批 vs 基线批 的 recommend 质量对比
print("\n=== 6) 增强批 vs 基线批 recommend 样例 ===")
def sample(flag, n=3):
    out=[]
    for e in DB:
        if (e['serial'] in proc) == flag:
            out.append((e['serial'], e.get('name'), e.get('recommend')))
        if len(out)>=n: break
    return out
print("  -- 增强批 --")
for s,nm,r in sample(True,4): print(f"    {s} {nm}: {r}")
print("  -- 基线批 --")
for s,nm,r in sample(False,4): print(f"    {s} {nm}: {r}")
